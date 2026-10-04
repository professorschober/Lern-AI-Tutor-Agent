from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
from pypdf import PdfReader
from docx import Document

MAX_BYTES = 20 * 1024 * 1024


def extract(name: str, content: bytes):
    suffix = Path(name).suffix.lower()
    if suffix not in {'.pdf', '.docx', '.md'}:
        raise ValueError('Bitte PDF, DOCX oder Markdown hochladen.')
    if len(content) > MAX_BYTES:
        raise ValueError('Die Datei überschreitet 20 MB.')
    try:
        if suffix == '.pdf':
            reader = PdfReader(BytesIO(content))
            if len(reader.pages) > 500:
                raise ValueError('Maximal 500 Seiten pro Dokument.')
            sections = [{'location': f'Seite {i + 1}', 'text': p.extract_text() or ''}
                        for i, p in enumerate(reader.pages)]
        elif suffix == '.docx':
            with ZipFile(BytesIO(content)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 50 * 1024 * 1024:
                    raise ValueError('Entpacktes Dokument ist zu groß.')
            doc = Document(BytesIO(content))
            text = '\n'.join([p.text for p in doc.paragraphs] +
                             [' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows])
            sections = [{'location': 'Dokument', 'text': text}]
        else:
            sections = [{'location': 'Markdown', 'text': content.decode('utf-8-sig')}]
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('Datei beschädigt, verschlüsselt oder nicht lesbar.') from exc
    sections = [s for s in sections if s['text'].strip()]
    if not sections:
        raise ValueError('Kein verwertbarer Text. Gescannte PDFs benötigen OCR und werden nicht unterstützt.')
    if sum(len(s['text']) for s in sections) > 1_000_000:
        raise ValueError('Extrahierter Text ist zu groß.')
    return sections
